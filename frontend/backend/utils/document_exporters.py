from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from fpdf import FPDF

from backend.utils.text_utils import sanitize_text


# =========================================================
# TEXT / PDF SAFE TEXT
# =========================================================

def sanitize_pdf_text(text: str) -> str:
    """
    Convert text to characters supported by the built-in
    FPDF Times font.
    """
    text = sanitize_text(text)

    replacements = {
        "₹": "Rs.",
        "€": "EUR ",
        "£": "GBP ",
        "—": "-",
        "–": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text


# =========================================================
# TXT EXPORT
# =========================================================

def format_txt(text: str) -> bytes:
    """
    Convert generated document text into a TXT file.
    """
    clean_text = sanitize_text(text)

    return clean_text.encode(
        "utf-8"
    )


# =========================================================
# DOCX EXPORT
# =========================================================

def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_path: str = ""
) -> bytes:
    """
    Create a professional DOCX legal document.
    """

    document = Document()

    # -----------------------------------------------------
    # PAGE SETTINGS
    # -----------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # -----------------------------------------------------
    # LOGO
    # -----------------------------------------------------

    if logo_path:

        try:

            if logo_path.lower().endswith(
                (
                    ".png",
                    ".jpg",
                    ".jpeg",
                    ".bmp"
                )
            ):

                paragraph = document.add_paragraph()

                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

                paragraph.add_run().add_picture(
                    logo_path,
                    width=Inches(1.5)
                )

        except Exception:
            # If the logo cannot be loaded,
            # continue without it.
            pass

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        sanitize_text(doc_type).upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(18)

    # -----------------------------------------------------
    # SUBTITLE
    # -----------------------------------------------------

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle_run = subtitle.add_run(
        "AI-GENERATED LEGAL DOCUMENT DRAFT"
    )

    subtitle_run.italic = True
    subtitle_run.font.name = "Times New Roman"
    subtitle_run.font.size = Pt(9)

    # -----------------------------------------------------
    # DOCUMENT CONTENT
    # -----------------------------------------------------

    clean_text = sanitize_text(text)

    paragraphs = clean_text.split("\n")

    for line in paragraphs:

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        run = paragraph.add_run(line)

        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    # -----------------------------------------------------
    # KEY TERMS
    # -----------------------------------------------------

    if terms.strip():

        document.add_paragraph()

        heading = document.add_paragraph()

        heading_run = heading.add_run(
            "KEY TERMS"
        )

        heading_run.bold = True
        heading_run.font.name = "Times New Roman"
        heading_run.font.size = Pt(13)

        table = document.add_table(
            rows=1,
            cols=1
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.style = "Table Grid"

        cell = table.rows[0].cells[0]

        cell.text = sanitize_text(
            terms
        )

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "Generated by LegalEase AI"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    # -----------------------------------------------------
    # SAVE TO MEMORY
    # -----------------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# =========================================================
# PDF CLASS
# =========================================================

class LegalEasePDF(FPDF):

    def __init__(self, document_type: str):

        super().__init__()

        self.document_type = (
            sanitize_pdf_text(document_type)
        )

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    def header(self):

        self.set_font(
            "Times",
            "B",
            15
        )

        self.cell(
            0,
            10,
            self.document_type.upper(),
            align="C"
        )

        self.ln(12)

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Times",
            "",
            8
        )

        self.cell(
            0,
            10,
            "Generated by LegalEase AI",
            align="C"
        )


# =========================================================
# PDF EXPORT
# =========================================================

def format_pdf(
    text: str,
    doc_type: str,
    terms: str = ""
) -> bytes:
    """
    Create a PDF legal document.

    The built-in FPDF Times font does not support
    characters such as the Indian Rupee symbol,
    so those are safely converted to Rs.
    """

    pdf = LegalEasePDF(
        doc_type
    )

    pdf.add_page()

    # -----------------------------------------------------
    # MAIN CONTENT FONT
    # -----------------------------------------------------

    pdf.set_font(
        "Times",
        "",
        11
    )

    clean_text = sanitize_pdf_text(
        text
    )

    paragraphs = clean_text.split(
        "\n"
    )

    for paragraph in paragraphs:

        paragraph = paragraph.strip()

        if not paragraph:

            pdf.ln(4)

            continue

        pdf.multi_cell(
            0,
            7,
            paragraph
        )

        pdf.ln(2)

    # -----------------------------------------------------
    # KEY TERMS
    # -----------------------------------------------------

    if terms.strip():

        pdf.ln(5)

        pdf.set_font(
            "Times",
            "B",
            13
        )

        pdf.cell(
            0,
            8,
            "KEY TERMS"
        )

        pdf.ln(9)

        pdf.set_font(
            "Times",
            "",
            10
        )

        pdf.multi_cell(
            0,
            6,
            sanitize_pdf_text(
                terms
            )
        )

    # -----------------------------------------------------
    # RETURN PDF BYTES
    # -----------------------------------------------------

    output = pdf.output()

    if isinstance(
        output,
        bytearray
    ):
        return bytes(output)

    return output